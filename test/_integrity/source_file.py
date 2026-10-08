from pathlib import Path
import sys
import csv
from datetime import datetime


# ====================================================================
# FORMAT / VALIDATOR INFORMATION
# ====================================================================

FORMAT_NAMES = {
    ".csv": "CSV",
    ".tsv": "TSV",
    ".txt": "TXT",
    ".json": "JSON",
    ".jsonl": "JSONL",
    ".xlsx": "Excel",
    ".parquet": "Parquet",
}


VALIDATOR_NAMES = {
    ".csv": "CSV Structure Validator",
    ".tsv": "TSV Structure Validator",
    ".txt": "TXT Structure Validator",
    ".json": "JSON Structure Validator",
    ".jsonl": "JSONL Structure Validator",
    ".xlsx": "Excel Structure Validator",
    ".parquet": "Parquet Structure Validator",
}


# ====================================================================
# HELPER
# ====================================================================

def _format_size(size_bytes: int) -> str:
    """
    Memformat ukuran bytes menjadi format yang mudah dibaca.
    """
    size = float(size_bytes)

    for unit in ["B", "KB", "MB", "GB"]:
        if size < 1024.0:
            return f"{size:.2f} {unit}"

        size /= 1024.0

    return f"{size:.2f} TB"


def _format_timestamp(timestamp: float) -> str:
    """
    Memformat timestamp filesystem menjadi tanggal yang mudah dibaca.
    """
    return datetime.fromtimestamp(timestamp).strftime(
        "%Y-%m-%d %H:%M:%S"
    )


def _detect_delimiter(file_path: Path) -> str:
    """
    Mendeteksi delimiter file text menggunakan csv.Sniffer.

    Jika gagal, gunakan koma sebagai default.
    """

    try:
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            sample = file.read(4096)

            if not sample.strip():
                return ","

            dialect = csv.Sniffer().sniff(
                sample,
                delimiters=",;\t|"
            )

            return dialect.delimiter

    except Exception:
        return ","


def _read_text_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar file text/tabular.

    Informasi:
        - encoding
        - delimiter
        - has_header
        - columns
        - column_count
        - row_count
        - empty_rows
        - inconsistent_rows
    """

    delimiter = _detect_delimiter(file_path)

    row_count = 0
    empty_rows = 0
    inconsistent_rows = 0
    columns = []

    try:
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.reader(
                file,
                delimiter=delimiter
            )

            try:
                header = next(reader)

            except StopIteration:
                return {
                    "encoding": "UTF-8",
                    "delimiter": delimiter,
                    "has_header": False,
                    "columns": [],
                    "column_count": 0,
                    "row_count": 0,
                    "empty_rows": 0,
                    "inconsistent_rows": 0,
                }

            columns = [
                column.strip()
                for column in header
            ]

            expected_column_count = len(columns)

            for row in reader:

                if not row or all(
                    str(value).strip() == ""
                    for value in row
                ):
                    empty_rows += 1
                    continue

                row_count += 1

                if len(row) != expected_column_count:
                    inconsistent_rows += 1

        return {
            "encoding": "UTF-8",
            "delimiter": delimiter,
            "has_header": True,
            "columns": columns,
            "column_count": len(columns),
            "row_count": row_count,
            "empty_rows": empty_rows,
            "inconsistent_rows": inconsistent_rows,
        }

    except UnicodeDecodeError:

        return {
            "encoding": "Unknown",
            "delimiter": delimiter,
            "has_header": False,
            "columns": [],
            "column_count": 0,
            "row_count": 0,
            "empty_rows": 0,
            "inconsistent_rows": 0,
        }

    except Exception:
        raise


def _read_json_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar JSON / JSONL.
    """

    import json

    row_count = 0
    empty_rows = 0
    columns = []
    inconsistent_rows = 0

    with open(
        file_path,
        "r",
        encoding="utf-8-sig"
    ) as file:

        content = file.read().strip()

    if not content:

        return {
            "encoding": "UTF-8",
            "delimiter": None,
            "has_header": False,
            "columns": [],
            "column_count": 0,
            "row_count": 0,
            "empty_rows": 0,
            "inconsistent_rows": 0,
        }

    # ================================================================
    # JSONL
    # ================================================================

    if file_path.suffix.lower() == ".jsonl":

        records = []

        for line in content.splitlines():

            if not line.strip():
                empty_rows += 1
                continue

            record = json.loads(line)
            records.append(record)

        if records and all(
            isinstance(record, dict)
            for record in records
        ):

            for record in records:
                columns.extend(record.keys())

            columns = list(dict.fromkeys(columns))
            row_count = len(records)

            return {
                "encoding": "UTF-8",
                "delimiter": None,
                "has_header": True,
                "columns": columns,
                "column_count": len(columns),
                "row_count": row_count,
                "empty_rows": empty_rows,
                "inconsistent_rows": inconsistent_rows,
            }

        return {
            "encoding": "UTF-8",
            "delimiter": None,
            "has_header": False,
            "columns": [],
            "column_count": 0,
            "row_count": len(records),
            "empty_rows": empty_rows,
            "inconsistent_rows": 0,
        }

    # ================================================================
    # JSON
    # ================================================================

    data = json.loads(content)

    # ================================================================
    # JSON ARRAY
    # ================================================================

    if isinstance(data, list):

        row_count = len(data)

        # ------------------------------------------------------------
        # JSON ARRAY OF OBJECTS
        # ------------------------------------------------------------

        if data and all(
            isinstance(record, dict)
            for record in data
        ):

            for record in data:
                columns.extend(record.keys())

            columns = list(dict.fromkeys(columns))

            return {
                "encoding": "UTF-8",
                "delimiter": None,
                "has_header": True,
                "columns": columns,
                "column_count": len(columns),
                "row_count": row_count,
                "empty_rows": 0,
                "inconsistent_rows": 0,
            }

        # ------------------------------------------------------------
        # JSON ARRAY OF ARRAYS
        # ------------------------------------------------------------

        if data and all(
            isinstance(record, list)
            for record in data
        ):

            columns = [
                f"column_{index + 1}"
                for index in range(
                    max(len(record) for record in data)
                )
            ]

            expected_column_count = len(columns)

            for record in data:

                if len(record) != expected_column_count:
                    inconsistent_rows += 1

            return {
                "encoding": "UTF-8",
                "delimiter": None,
                "has_header": False,
                "columns": columns,
                "column_count": len(columns),
                "row_count": row_count,
                "empty_rows": 0,
                "inconsistent_rows": inconsistent_rows,
            }

    # ================================================================
    # JSON OBJECT
    # ================================================================

    if isinstance(data, dict):

        columns = list(data.keys())

        return {
            "encoding": "UTF-8",
            "delimiter": None,
            "has_header": True,
            "columns": columns,
            "column_count": len(columns),
            "row_count": 1,
            "empty_rows": 0,
            "inconsistent_rows": 0,
        }

    # ================================================================
    # JSON SCALAR
    # ================================================================

    return {
        "encoding": "UTF-8",
        "delimiter": None,
        "has_header": False,
        "columns": [],
        "column_count": 0,
        "row_count": 1,
        "empty_rows": 0,
        "inconsistent_rows": 0,
    }


def _read_excel_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar Excel.

    Membutuhkan openpyxl untuk file .xlsx.
    """

    from openpyxl import load_workbook

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    try:

        worksheet = workbook.active

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            header = next(rows)

        except StopIteration:

            return {
                "encoding": "Excel",
                "delimiter": None,
                "has_header": False,
                "columns": [],
                "column_count": 0,
                "row_count": 0,
                "empty_rows": 0,
                "inconsistent_rows": 0,
            }

        columns = [
            str(column).strip()
            if column is not None
            else ""
            for column in header
        ]

        expected_column_count = len(columns)

        row_count = 0
        empty_rows = 0
        inconsistent_rows = 0

        for row in rows:

            if not row or all(
                value is None or str(value).strip() == ""
                for value in row
            ):
                empty_rows += 1
                continue

            row_count += 1

            if len(row) != expected_column_count:
                inconsistent_rows += 1

        return {
            "encoding": "Excel",
            "delimiter": None,
            "has_header": True,
            "columns": columns,
            "column_count": len(columns),
            "row_count": row_count,
            "empty_rows": empty_rows,
            "inconsistent_rows": inconsistent_rows,
        }

    finally:
        workbook.close()


def _read_parquet_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar Parquet.

    Membutuhkan pyarrow.
    """

    import pyarrow.parquet as pq

    parquet_file = pq.ParquetFile(
        file_path
    )

    schema = parquet_file.schema_arrow

    columns = list(schema.names)

    return {
        "encoding": "Parquet",
        "delimiter": None,
        "has_header": True,
        "columns": columns,
        "column_count": len(columns),
        "row_count": parquet_file.metadata.num_rows,
        "empty_rows": 0,
        "inconsistent_rows": 0,
    }


def _read_csv_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar CSV.
    """

    return _read_text_metadata(file_path)


def _read_tsv_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar TSV.
    """

    return _read_text_metadata(file_path)


def _read_txt_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar TXT.
    """

    return _read_text_metadata(file_path)


def _read_jsonl_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar JSONL.
    """

    return _read_json_metadata(file_path)


def _read_file_metadata(file_path: Path) -> dict:
    """
    Membaca metadata dasar source file berdasarkan extension.

    Format yang didukung:
        - CSV
        - TSV
        - TXT
        - JSON
        - JSONL
        - XLSX
        - Parquet
    """

    extension = file_path.suffix.lower()

    if extension == ".csv":

        return _read_csv_metadata(
            file_path
        )

    elif extension == ".tsv":

        return _read_tsv_metadata(
            file_path
        )

    elif extension == ".txt":

        return _read_txt_metadata(
            file_path
        )

    elif extension == ".json":

        return _read_json_metadata(
            file_path
        )

    elif extension == ".jsonl":

        return _read_jsonl_metadata(
            file_path
        )

    elif extension == ".xlsx":

        return _read_excel_metadata(
            file_path
        )

    elif extension == ".parquet":

        return _read_parquet_metadata(
            file_path
        )

    else:

        raise ValueError(
            f"Format file '{extension or '(none)'}' "
            f"belum didukung."
        )


# ====================================================================
# SUPPORTED EXTENSIONS
# ====================================================================

SUPPORTED_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".txt",
    ".json",
    ".jsonl",
    ".xlsx",
    ".parquet",
}


# ====================================================================
# MAIN SOURCE FILE VALIDATION
# ====================================================================

def check_source_file(
    file_path: Path,
    dataset_name: str = "dataset"
) -> dict:
    """
    Memvalidasi source file data.

    Validation:
        1. Check Existence
        2. Check Data
        3. Check Extension / Format
        4. Check Integrity
        5. Check Structure
    """

    # ================================================================
    # HEADER
    # ================================================================

    print("\n" + "=" * 70)

    print(
        f"[INFO] Memulai Pengecekan Dataset: "
        f"'{dataset_name}'"
    )

    print("=" * 70)

    print()

    print("[DEBUG] Target:")

    print(
        f"  ├─ Path      : "
        f"{file_path.resolve()}"
    )

    print(
        f"  ├─ File      : "
        f"{file_path.name}"
    )

    print(
        f"  └─ Dataset   : "
        f"{dataset_name}"
    )

    # ================================================================
    # [1/5] CHECK EXISTENCE
    # ================================================================

    try:

        print()

        print(
            "[DEBUG] Check Existence"
        )

        # ------------------------------------------------------------
        # pathlib.exists() akan False untuk broken symlink.
        #
        # Karena itu kita juga mengecek is_symlink().
        #
        # Dengan cara ini broken symlink tetap dianggap sebagai
        # filesystem entry dan bisa diperiksa secara eksplisit
        # pada tahap [2/5] Check Data.
        # ------------------------------------------------------------

        exists = file_path.exists()
        is_symlink_entry = file_path.is_symlink()

        path_entry_exists = (
            exists or is_symlink_entry
        )

        print(
            "  ├─ Expected  : "
            "File/path entry exists"
        )

        print(
            f"  ├─ Actual    : "
            f"{path_entry_exists}"
        )

        if is_symlink_entry:

            print(
                "  ├─ Note      : "
                "Path merupakan symbolic link"
            )

        else:

            print(
                "  ├─ Note      : "
                "Path merupakan filesystem entry biasa"
            )

        if not path_entry_exists:

            print(
                "  └─ Result    : "
                "FAIL"
            )

            print()

            print(
                f"[FAIL] Tahap 1 Gagal -> "
                f"Path '{dataset_name}' "
                f"tidak ditemukan."
            )

            print("-" * 70)

            return {
                "status": "FAIL",
                "actual": (
                    f"{file_path} does not exist"
                ),
                "expected": (
                    "file/path entry exists"
                ),
                "message": (
                    f"Source file '{dataset_name}' "
                    f"tidak ditemukan: "
                    f"{file_path}"
                )
            }

        print(
            "  └─ Result    : "
            "PASS"
        )

    except Exception as e:

        print(
            "  ├─ Expected  : "
            "File/path entry exists"
        )

        print(
            f"  ├─ Actual    : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result    : "
            "ERROR"
        )

        print(
            f"[ERROR] Tahap 1 Exception -> "
            f"{str(e)}",
            file=sys.stderr
        )

        return {
            "status": "ERROR",
            "actual": type(e).__name__,
            "expected": (
                "successful path existence check"
            ),
            "message": (
                f"Gagal mengecek keberadaan file "
                f"{file_path}: {str(e)}"
            )
        }

    # ================================================================
    # [2/5] CHECK DATA
    # ================================================================

    try:

        print()

        print(
            "[DEBUG] Check Data"
        )

        is_file = file_path.is_file()
        is_directory = file_path.is_dir()
        is_symlink = file_path.is_symlink()

        # ------------------------------------------------------------
        # DEFAULT
        # ------------------------------------------------------------

        path_type = "Unknown"
        symlink_target = None
        symlink_target_exists = None
        symlink_target_is_file = None

        # ------------------------------------------------------------
        # SYMLINK ANALYSIS
        # ------------------------------------------------------------

        if is_symlink:

            try:

                symlink_target = file_path.resolve(
                    strict=False
                )

                symlink_target_exists = (
                    symlink_target.exists()
                )

                symlink_target_is_file = (
                    symlink_target.is_file()
                )

                if (
                    symlink_target_exists
                    and symlink_target_is_file
                ):

                    path_type = (
                        "Symbolic Link → File"
                    )

                elif (
                    symlink_target_exists
                    and symlink_target.is_dir()
                ):

                    path_type = (
                        "Symbolic Link → Directory"
                    )

                else:

                    path_type = (
                        "Broken Symbolic Link"
                    )

            except OSError as error:

                path_type = (
                    "Symbolic Link → Unknown"
                )

                symlink_target = None

                print(
                    "  ├─ Symlink Error : "
                    f"{error}"
                )

        else:

            # --------------------------------------------------------
            # REGULAR FILE / DIRECTORY
            # --------------------------------------------------------

            if is_file:

                path_type = "Regular File"

            elif is_directory:

                path_type = "Directory"

            else:

                path_type = "Unknown"

        # ------------------------------------------------------------
        # OUTPUT
        # ------------------------------------------------------------

        print(
            "  ├─ Expected      : "
            "Regular data file"
        )

        print(
            f"  ├─ Path Type     : "
            f"{path_type}"
        )

        print(
            f"  ├─ Is File       : "
            f"{is_file}"
        )

        print(
            f"  ├─ Is Directory  : "
            f"{is_directory}"
        )

        print(
            f"  ├─ Is Symlink    : "
            f"{is_symlink}"
        )

        # ------------------------------------------------------------
        # SYMLINK DETAILS
        # ------------------------------------------------------------

        if is_symlink:

            if (
                symlink_target_exists
                and symlink_target_is_file
            ):

                print(
                    "  ├─ Symlink Info  : "
                    "Symbolic link valid "
                    "dan target berupa file"
                )

                print(
                    f"  ├─ Link Target   : "
                    f"{symlink_target}"
                )

                print(
                    "  ├─ Target Exists : "
                    f"{symlink_target_exists}"
                )

                print(
                    "  ├─ Target File   : "
                    f"{symlink_target_is_file}"
                )

            elif (
                symlink_target_exists
                and symlink_target.is_dir()
            ):

                print(
                    "  ├─ Symlink Info  : "
                    "Symbolic link valid, "
                    "tetapi target berupa directory"
                )

                print(
                    f"  ├─ Link Target   : "
                    f"{symlink_target}"
                )

                print(
                    "  ├─ Target Exists : "
                    f"{symlink_target_exists}"
                )

                print(
                    "  ├─ Target File   : "
                    f"{symlink_target_is_file}"
                )

            else:

                print(
                    "  ├─ Symlink Info  : "
                    "BROKEN SYMBOLIC LINK - "
                    "target tidak ditemukan"
                )

                print(
                    f"  ├─ Link Target   : "
                    f"{symlink_target}"
                )

                print(
                    "  ├─ Target Exists : "
                    f"{symlink_target_exists}"
                )

                print(
                    "  ├─ Target File   : "
                    f"{symlink_target_is_file}"
                )

        else:

            print(
                "  ├─ Symlink Info  : "
                "Bukan symbolic link; "
                "path langsung menunjuk ke file"
            )

        # ------------------------------------------------------------
        # VALIDATION
        # ------------------------------------------------------------

        if (
            is_file
            and not is_directory
            and (
                not is_symlink
                or (
                    symlink_target_exists
                    and symlink_target_is_file
                )
            )
        ):

            print(
                "  └─ Result        : "
                "PASS"
            )

        else:

            print(
                "  └─ Result        : "
                "FAIL"
            )

            print()

            if is_symlink:

                if not symlink_target_exists:

                    message = (
                        f"Symbolic link '{dataset_name}' "
                        f"merupakan broken symlink. "
                        f"Target tidak ditemukan: "
                        f"{symlink_target}"
                    )

                elif symlink_target.is_dir():

                    message = (
                        f"Symbolic link '{dataset_name}' "
                        f"mengarah ke directory, "
                        f"bukan file."
                    )

                else:

                    message = (
                        f"Symbolic link '{dataset_name}' "
                        f"tidak mengarah ke regular file."
                    )

            elif is_directory:

                message = (
                    f"Path '{dataset_name}' "
                    f"merupakan directory, "
                    f"bukan regular file."
                )

            else:

                message = (
                    f"Path '{dataset_name}' "
                    f"bukan regular file."
                )

            print(
                f"[FAIL] Tahap 2 Gagal -> "
                f"{message}"
            )

            print("-" * 70)

            return {
                "status": "FAIL",
                "actual": {
                    "is_file": is_file,
                    "is_directory": is_directory,
                    "is_symlink": is_symlink,
                    "path_type": path_type,
                    "symlink_target": (
                        str(symlink_target)
                        if symlink_target
                        else None
                    ),
                    "symlink_target_exists": (
                        symlink_target_exists
                    ),
                },
                "expected": (
                    "regular data file"
                ),
                "message": message
            }

    except Exception as e:

        print(
            "  ├─ Expected      : "
            "Regular data file"
        )

        print(
            f"  ├─ Actual        : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result        : "
            "ERROR"
        )

        print(
            f"[ERROR] Tahap 2 Exception -> "
            f"{str(e)}",
            file=sys.stderr
        )

        return {
            "status": "ERROR",
            "actual": type(e).__name__,
            "expected": (
                "successful data type check"
            ),
            "message": (
                f"Gagal mengecek tipe data "
                f"{file_path}: {str(e)}"
            )
        }

    # ================================================================
    # [3/5] CHECK EXTENSION / FORMAT
    # ================================================================

    try:

        print()

        print(
            "[DEBUG] Check Extension / Format"
        )

        ext = file_path.suffix.lower()
        stem = file_path.stem
        filename = file_path.name

        format_name = FORMAT_NAMES.get(
            ext,
            "Unknown"
        )

        validator_name = VALIDATOR_NAMES.get(
            ext,
            "Unknown Validator"
        )

        print(
            "  ├─ Expected  : "
            "Supported data file"
        )

        print(
            f"  ├─ File Name : "
            f"{filename}"
        )

        print(
            f"  ├─ File Stem : "
            f"{stem}"
        )

        print(
            f"  ├─ Extension : "
            f"{ext or '(none)'}"
        )

        print(
            f"  ├─ Format    : "
            f"{format_name}"
        )

        print(
            f"  ├─ Validator : "
            f"{validator_name}"
        )

        print(
            "  ├─ Strategy  : "
            f"Pemeriksaan file akan disesuaikan "
            f"dengan format {format_name}"
        )

        print(
            f"  ├─ Supported : "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

        if ext not in SUPPORTED_EXTENSIONS:

            print(
                "  └─ Result    : "
                "FAIL"
            )

            print()

            print(
                f"[FAIL] Tahap 3 Gagal -> "
                f"Format file '{dataset_name}' "
                f"tidak didukung."
            )

            print("-" * 70)

            return {
                "status": "FAIL",
                "actual": ext,
                "expected": sorted(
                    SUPPORTED_EXTENSIONS
                ),
                "message": (
                    f"Source file '{dataset_name}' "
                    f"menggunakan format yang "
                    f"belum didukung: "
                    f"{file_path}"
                )
            }

        print(
            "  └─ Result    : "
            "PASS"
        )

    except Exception as e:

        print(
            "  ├─ Expected  : "
            "Supported data file"
        )

        print(
            f"  ├─ Actual    : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result    : "
            "ERROR"
        )

        print(
            f"[ERROR] Tahap 3 Exception -> "
            f"{str(e)}",
            file=sys.stderr
        )

        return {
            "status": "ERROR",
            "actual": type(e).__name__,
            "expected": (
                "successful extension check"
            ),
            "message": (
                f"Gagal membaca ekstensi file "
                f"{file_path}: {str(e)}"
            )
        }

    # ================================================================
    # [4/5] CHECK INTEGRITY
    # ================================================================

    try:

        print()

        print(
            "[DEBUG] Check Integrity"
        )

        file_stat = file_path.stat()

        size_bytes = file_stat.st_size

        readable_size = _format_size(
            size_bytes
        )

        modified = _format_timestamp(
            file_stat.st_mtime
        )

        created = _format_timestamp(
            file_stat.st_ctime
        )

        accessed = _format_timestamp(
            file_stat.st_atime
        )

        non_empty = (
            size_bytes > 0
        )

        can_read = False
        can_write = False

        # ------------------------------------------------------------
        # READ TEST
        # ------------------------------------------------------------

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8-sig"
            ):

                can_read = True

        except Exception:

            can_read = False

        # ------------------------------------------------------------
        # WRITE TEST
        # ------------------------------------------------------------

        try:

            with open(
                file_path,
                "a",
                encoding="utf-8"
            ):

                can_write = True

        except Exception:

            can_write = False

        # ------------------------------------------------------------
        # OUTPUT
        # ------------------------------------------------------------

        print(
            f"  ├─ Size Bytes    : "
            f"{size_bytes:,} Bytes"
        )

        print(
            f"  ├─ Size Human    : "
            f"{readable_size}"
        )

        print(
            f"  ├─ Non Empty     : "
            f"{non_empty}"
        )

        print(
            f"  ├─ Readable      : "
            f"{can_read}"
        )

        print(
            f"  ├─ Writable      : "
            f"{can_write}"
        )

        print(
            f"  ├─ Created       : "
            f"{created}"
        )

        print(
            f"  ├─ Modified      : "
            f"{modified}"
        )

        print(
            f"  ├─ Last Access   : "
            f"{accessed}"
        )

        # ------------------------------------------------------------
        # EMPTY FILE
        # ------------------------------------------------------------

        if not non_empty:

            print(
                "  └─ Result        : "
                "FAIL"
            )

            print()

            print(
                f"[FAIL] Tahap 4 Gagal -> "
                f"File '{dataset_name}' kosong."
            )

            print("-" * 70)

            return {
                "status": "FAIL",
                "actual": size_bytes,
                "expected": "> 0 bytes",
                "message": (
                    f"Source file '{dataset_name}' "
                    f"kosong: {file_path}"
                )
            }

        # ------------------------------------------------------------
        # READABLE CHECK
        # ------------------------------------------------------------

        if not can_read:

            print(
                "  └─ Result        : "
                "FAIL"
            )

            print()

            print(
                f"[FAIL] Tahap 4 Gagal -> "
                f"File '{dataset_name}' "
                f"tidak dapat dibaca."
            )

            print("-" * 70)

            return {
                "status": "FAIL",
                "actual": "file not readable",
                "expected": "readable file",
                "message": (
                    f"Source file '{dataset_name}' "
                    f"tidak dapat dibaca."
                )
            }

        print(
            "  └─ Result        : "
            "PASS"
        )

    except Exception as e:

        print(
            "  ├─ Expected  : "
            "Valid readable file"
        )

        print(
            f"  ├─ Actual    : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result    : "
            "ERROR"
        )

        print(
            f"[ERROR] Tahap 4 Exception -> "
            f"{str(e)}",
            file=sys.stderr
        )

        return {
            "status": "ERROR",
            "actual": type(e).__name__,
            "expected": (
                "successful file integrity check"
            ),
            "message": (
                f"Gagal membaca metadata file "
                f"{file_path}: {str(e)}"
            )
        }

    # ================================================================
    # [5/5] CHECK STRUCTURE
    # ================================================================

    try:

        print()

        print(
            "[DEBUG] Check Structure"
        )

        # ------------------------------------------------------------
        # FORMAT / VALIDATOR
        # ------------------------------------------------------------

        format_name = FORMAT_NAMES.get(
            ext,
            "Unknown"
        )

        validator_name = VALIDATOR_NAMES.get(
            ext,
            "Unknown Validator"
        )

        print(
            f"  ├─ Format          : "
            f"{format_name}"
        )

        print(
            f"  ├─ Validator       : "
            f"{validator_name}"
        )

        print(
            "  ├─ Strategy        : "
            f"Pemeriksaan struktur "
            f"berdasarkan format {format_name}"
        )

        # ------------------------------------------------------------
        # READ METADATA
        # ------------------------------------------------------------

        file_metadata = _read_file_metadata(
            file_path
        )

        delimiter = file_metadata[
            "delimiter"
        ]

        has_header = file_metadata[
            "has_header"
        ]

        columns = file_metadata[
            "columns"
        ]

        column_count = file_metadata[
            "column_count"
        ]

        row_count = file_metadata[
            "row_count"
        ]

        empty_rows = file_metadata[
            "empty_rows"
        ]

        inconsistent_rows = file_metadata[
            "inconsistent_rows"
        ]

        # ------------------------------------------------------------
        # ENCODING
        # ------------------------------------------------------------

        print(
            f"  ├─ Encoding        : "
            f"{file_metadata['encoding']}"
        )

        # ------------------------------------------------------------
        # DELIMITER
        # ------------------------------------------------------------

        delimiter_display = {
            ",": "Comma (,)",
            ";": "Semicolon (;)",
            "\t": "Tab (\\t)",
            "|": "Pipe (|)"
        }.get(
            delimiter,
            repr(delimiter)
        )

        print(
            f"  ├─ Delimiter       : "
            f"{delimiter_display}"
        )

        # ------------------------------------------------------------
        # HEADER
        # ------------------------------------------------------------

        print(
            f"  ├─ Header          : "
            f"{has_header}"
        )

        # ------------------------------------------------------------
        # COLUMNS
        # ------------------------------------------------------------

        print(
            f"  ├─ Column Count    : "
            f"{column_count}"
        )

        # ------------------------------------------------------------
        # ROWS
        # ------------------------------------------------------------

        print(
            f"  ├─ Data Rows       : "
            f"{row_count:,}"
        )

        print(
            f"  ├─ Empty Rows      : "
            f"{empty_rows:,}"
        )

        print(
            f"  ├─ Inconsistent    : "
            f"{inconsistent_rows:,}"
        )

        # ------------------------------------------------------------
        # COLUMN LIST
        # ------------------------------------------------------------

        print()

        print(
            "  ├─ Columns:"
        )

        if columns:

            for index, column in enumerate(
                columns,
                start=1
            ):

                connector = (
                    "└─"
                    if index == len(columns)
                    else "├─"
                )

                print(
                    f"  │  {connector} "
                    f"{index:02d}. {column}"
                )

        else:

            print(
                "  │  └─ (no columns)"
            )

        # ------------------------------------------------------------
        # HEADER VALIDATION
        # ------------------------------------------------------------

        if not has_header:

            print()

            print(
                "  └─ Result          : "
                "FAIL"
            )

            return {
                "status": "FAIL",
                "actual": (
                    "Data file has no header"
                ),
                "expected": (
                    "Data file with header"
                ),
                "message": (
                    f"Data file '{dataset_name}' "
                    f"tidak memiliki header."
                )
            }

        # ------------------------------------------------------------
        # COLUMN VALIDATION
        # ------------------------------------------------------------

        if column_count == 0:

            print()

            print(
                "  └─ Result          : "
                "FAIL"
            )

            return {
                "status": "FAIL",
                "actual": 0,
                "expected": "> 0 columns",
                "message": (
                    f"Data file '{dataset_name}' "
                    f"tidak memiliki kolom."
                )
            }

        # ------------------------------------------------------------
        # CONSISTENCY VALIDATION
        # ------------------------------------------------------------

        if inconsistent_rows > 0:

            print()

            print(
                "  └─ Result          : "
                "FAIL"
            )

            return {
                "status": "FAIL",
                "actual": inconsistent_rows,
                "expected": 0,
                "message": (
                    f"Data file '{dataset_name}' "
                    f"memiliki {inconsistent_rows} "
                    f"baris dengan jumlah kolom "
                    f"tidak konsisten."
                )
            }

        print()

        print(
            "  └─ Result          : "
            "PASS"
        )

    except UnicodeDecodeError as e:

        print(
            "  ├─ Expected        : "
            "Valid data file"
        )

        print(
            f"  ├─ Actual          : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result          : "
            "FAIL"
        )

        return {
            "status": "FAIL",
            "actual": "invalid encoding",
            "expected": "valid data file",
            "message": (
                f"Data file '{dataset_name}' "
                f"tidak menggunakan encoding UTF-8."
            )
        }

    except Exception as e:

        print(
            "  ├─ Expected        : "
            "Valid data structure"
        )

        print(
            f"  ├─ Actual          : "
            f"{type(e).__name__}"
        )

        print(
            "  └─ Result          : "
            "ERROR"
        )

        print(
            f"[ERROR] Tahap 5 Exception -> "
            f"{str(e)}",
            file=sys.stderr
        )

        return {
            "status": "ERROR",
            "actual": type(e).__name__,
            "expected": (
                "successful data structure check"
            ),
            "message": (
                f"Gagal membaca struktur data "
                f"{file_path}: {str(e)}"
            )
        }

    # ================================================================
    # FINAL RESULT
    # ================================================================

    print()

    print("-" * 70)

    print(
        f"[PASS] VALIDASI SUKSES: "
        f"File '{dataset_name}' memenuhi "
        f"seluruh kriteria."
    )

    print("-" * 70)

    # ================================================================
    # VALIDATION SUMMARY
    # ================================================================

    print()

    print(
        "[DEBUG] Validation Summary:"
    )

    # ------------------------------------------------------------
    # DATASET
    # ------------------------------------------------------------

    print(
        f"  ├─ Dataset          : "
        f"{dataset_name}"
    )

    # ------------------------------------------------------------
    # FILE
    # ------------------------------------------------------------

    print(
        f"  ├─ File             : "
        f"{file_path.name}"
    )

    # ------------------------------------------------------------
    # ABSOLUTE PATH
    # ------------------------------------------------------------

    print(
        f"  ├─ Absolute Path    : "
        f"{file_path.resolve()}"
    )

    # ------------------------------------------------------------
    # PATH TYPE
    # ------------------------------------------------------------

    print(
        f"  ├─ Type             : "
        f"{format_name} / {path_type}"
    )

    # ------------------------------------------------------------
    # SYMLINK
    # ------------------------------------------------------------

    print(
        f"  ├─ Is Symlink       : "
        f"{is_symlink}"
    )

    if is_symlink:

        print(
            f"  ├─ Symlink Info     : "
            f"Symbolic link valid"
        )

        print(
            f"  ├─ Link Target      : "
            f"{symlink_target}"
        )

    else:

        print(
            f"  ├─ Symlink Info     : "
            f"Bukan symbolic link"
        )

    # ------------------------------------------------------------
    # EXTENSION
    # ------------------------------------------------------------

    print(
        f"  ├─ Extension        : "
        f"{ext}"
    )

    # ------------------------------------------------------------
    # FORMAT
    # ------------------------------------------------------------

    print(
        f"  ├─ Format           : "
        f"{format_name}"
    )

    # ------------------------------------------------------------
    # VALIDATOR
    # ------------------------------------------------------------

    print(
        f"  ├─ Validator        : "
        f"{validator_name}"
    )

    # ------------------------------------------------------------
    # SIZE
    # ------------------------------------------------------------

    print(
        f"  ├─ Size             : "
        f"{readable_size}"
    )

    print(
        f"  ├─ Size Bytes       : "
        f"{size_bytes:,}"
    )

    # ------------------------------------------------------------
    # ACCESS
    # ------------------------------------------------------------

    print(
        f"  ├─ Readable         : "
        f"{can_read}"
    )

    print(
        f"  ├─ Writable         : "
        f"{can_write}"
    )

    # ------------------------------------------------------------
    # TIMESTAMP
    # ------------------------------------------------------------

    print(
        f"  ├─ Created          : "
        f"{created}"
    )

    print(
        f"  ├─ Modified         : "
        f"{modified}"
    )

    print(
        f"  ├─ Last Access      : "
        f"{accessed}"
    )

    # ------------------------------------------------------------
    # STRUCTURE
    # ------------------------------------------------------------

    print(
        f"  ├─ Encoding         : "
        f"{file_metadata['encoding']}"
    )

    print(
        f"  ├─ Delimiter        : "
        f"{delimiter_display}"
    )

    print(
        f"  ├─ Header           : "
        f"{has_header}"
    )

    print(
        f"  ├─ Column Count     : "
        f"{column_count}"
    )

    print(
        f"  ├─ Data Rows        : "
        f"{row_count:,}"
    )

    print(
        f"  ├─ Empty Rows       : "
        f"{empty_rows:,}"
    )

    print(
        f"  ├─ Inconsistent     : "
        f"{inconsistent_rows:,}"
    )

    # ------------------------------------------------------------
    # FINAL
    # ------------------------------------------------------------

    print(
        f"  └─ Result           : "
        f"5/5 PASS"
    )

    print(
        "=" * 70
    )

    # ================================================================
    # RETURN
    # ================================================================

    return {
        "status": "PASS",

        "actual": {

            "path": str(
                file_path.resolve()
            ),

            "file_name": file_path.name,

            "extension": ext,

            "format": format_name,

            "validator": validator_name,

            "file": {

                "is_file": is_file,

                "is_directory": is_directory,

                "is_symlink": is_symlink,

                "path_type": path_type,

                "symlink_target": (
                    str(symlink_target)
                    if symlink_target
                    else None
                ),

                "symlink_target_exists": (
                    symlink_target_exists
                ),

                "symlink_target_is_file": (
                    symlink_target_is_file
                ),

                "readable": can_read,

                "writable": can_write,
            },

            "size": {

                "bytes": size_bytes,

                "human": readable_size,
            },

            "timestamp": {

                "created_at": created,

                "modified_at": modified,

                "last_access_at": accessed,
            },

            "data": {

                "encoding": (
                    file_metadata["encoding"]
                ),

                "delimiter": delimiter,

                "has_header": has_header,

                "column_count": column_count,

                "columns": columns,

                "row_count": row_count,

                "empty_rows": empty_rows,

                "inconsistent_rows": (
                    inconsistent_rows
                ),
            },
        },

        "expected": {

            "file_exists": True,

            "regular_file": True,

            "extension": sorted(
                SUPPORTED_EXTENSIONS
            ),

            "readable": True,

            "non_empty": True,

            "has_header": True,

            "column_count": "> 0",

            "inconsistent_rows": 0,
        },

        "message": (
            f"Source file '{dataset_name}' "
            f"valid."
        ),
    }