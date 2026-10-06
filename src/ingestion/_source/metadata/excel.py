# source_validator/metadata/excel.py

from pathlib import Path

from openpyxl import load_workbook


def read_excel_metadata(file_path: Path) -> dict:
    """
    Membaca metadata Excel.
    """

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    try:

        worksheet = workbook.active

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            header = next(rows)

        except StopIteration:
            return {
                "encoding": "Excel",
                "delimiter": None,
                "has_header": False,
                "columns": [],
                "column_count": 0,
                "row_count": 0,
                "empty_rows": 0,
                "inconsistent_rows": 0,
            }

        columns = [
            str(column).strip()
            if column is not None
            else ""
            for column in header
        ]

        expected_column_count = len(columns)

        row_count = 0
        empty_rows = 0
        inconsistent_rows = 0

        for row in rows:

            if not row or all(
                value is None
                or str(value).strip() == ""
                for value in row
            ):
                empty_rows += 1
                continue

            row_count += 1

            if len(row) != expected_column_count:
                inconsistent_rows += 1

        return {
            "encoding": "Excel",
            "delimiter": None,
            "has_header": True,
            "columns": columns,
            "column_count": len(columns),
            "row_count": row_count,
            "empty_rows": empty_rows,
            "inconsistent_rows": inconsistent_rows,
        }

    finally:
        workbook.close()